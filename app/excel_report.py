"""Write an AnalysisResult to a formatted Excel workbook."""
from __future__ import annotations

import io
from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .models import CATEGORIES, AnalysisResult, severity_for

SEVERITIES = ["Critical", "High", "Medium", "Low"]
SEVERITY_FILL = {
    "Critical": "C00000",
    "High": "FF7C43",
    "Medium": "FFD966",
    "Low": "A9D08E",
}
SEVERITY_FONT = {"Critical": "FFFFFF", "High": "000000", "Medium": "000000", "Low": "000000"}

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=16, color="1F3864")
BOLD = Font(bold=True)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP_TOP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _header(ws: Worksheet, row: int, headers: list[str], widths: list[int] | None = None) -> None:
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=col, value=h)
        c.fill, c.font, c.alignment, c.border = HEADER_FILL, HEADER_FONT, CENTER, BORDER
        if widths:
            ws.column_dimensions[get_column_letter(col)].width = widths[col - 1]


def _sev_cell(cell, severity: str) -> None:
    cell.fill = PatternFill("solid", fgColor=SEVERITY_FILL[severity])
    cell.font = Font(bold=True, color=SEVERITY_FONT[severity])
    cell.alignment = CENTER


def _summary(ws: Worksheet, result: AnalysisResult) -> None:
    ws.title = "Summary"
    ws.column_dimensions["A"].width = 30
    for col in "BCDEFG":
        ws.column_dimensions[col].width = 14

    ws["A1"] = f"Technical Risk Assessment – {result.project_name}"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:G1")

    p = result.profile
    info = [
        ("Domain", p.domain),
        ("Summary", p.summary),
        ("Key systems", ", ".join(p.key_systems)),
        ("Documents", ", ".join(f"{d.name} ({d.doc_type})" for d in result.documents)),
        ("Generated", result.finished_at),
        ("Total risks", len(result.risks)),
    ]
    row = 3
    for label, value in info:
        ws.cell(row=row, column=1, value=label).font = BOLD
        c = ws.cell(row=row, column=2, value=value)
        c.alignment = WRAP_TOP
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=7)
        if label == "Summary" and value:
            ws.row_dimensions[row].height = max(15, 15 * (len(str(value)) // 90 + 1))
        row += 1

    # Severity x category matrix
    row += 1
    ws.cell(row=row, column=1, value="Risks by category and severity").font = Font(bold=True, size=12)
    row += 1
    _header(ws, row, ["Category", *SEVERITIES, "Total"])
    counts = Counter((r.category, r.severity) for r in result.risks)
    for cat in CATEGORIES:
        row += 1
        ws.cell(row=row, column=1, value=cat).border = BORDER
        total = 0
        for i, sev in enumerate(SEVERITIES, 2):
            n = counts.get((cat, sev), 0)
            total += n
            c = ws.cell(row=row, column=i, value=n or None)
            c.alignment, c.border = CENTER, BORDER
            if n:
                _sev_cell(c, sev)
        c = ws.cell(row=row, column=6, value=total)
        c.font, c.alignment, c.border = BOLD, CENTER, BORDER
    row += 1
    ws.cell(row=row, column=1, value="Total").font = BOLD
    sev_totals = Counter(r.severity for r in result.risks)
    for i, sev in enumerate(SEVERITIES, 2):
        c = ws.cell(row=row, column=i, value=sev_totals.get(sev, 0))
        c.font, c.alignment, c.border = BOLD, CENTER, BORDER
    c = ws.cell(row=row, column=6, value=len(result.risks))
    c.font, c.alignment, c.border = BOLD, CENTER, BORDER

    # Heat map: likelihood (rows, 5 at top) x impact (cols)
    row += 2
    ws.cell(row=row, column=1, value="Heat map (number of risks)").font = Font(bold=True, size=12)
    row += 1
    ws.cell(row=row, column=1, value="Likelihood ↓  /  Impact →").font = BOLD
    for imp in range(1, 6):
        c = ws.cell(row=row, column=imp + 1, value=imp)
        c.fill, c.font, c.alignment = HEADER_FILL, HEADER_FONT, CENTER
    grid = Counter((r.likelihood, r.impact) for r in result.risks)
    for lik in range(5, 0, -1):
        row += 1
        c = ws.cell(row=row, column=1, value=lik)
        c.fill, c.font, c.alignment = HEADER_FILL, HEADER_FONT, CENTER
        for imp in range(1, 6):
            cell = ws.cell(row=row, column=imp + 1, value=grid.get((lik, imp)) or None)
            cell.border = BORDER
            sev = severity_for(lik * imp)
            cell.fill = PatternFill("solid", fgColor=SEVERITY_FILL[sev])
            cell.font = Font(bold=True, color=SEVERITY_FONT[sev])
            cell.alignment = CENTER

    # Top risks
    row += 2
    ws.cell(row=row, column=1, value="Top 10 risks").font = Font(bold=True, size=12)
    row += 1
    _header(ws, row, ["Title", "ID", "Severity", "Score", "Category"])
    ws.merge_cells(start_row=row, start_column=5, end_row=row, end_column=7)
    for r in result.risks[:10]:
        row += 1
        ws.cell(row=row, column=1, value=r.title).alignment = WRAP_TOP
        ws.cell(row=row, column=2, value=r.risk_id).alignment = CENTER
        _sev_cell(ws.cell(row=row, column=3, value=r.severity), r.severity)
        ws.cell(row=row, column=4, value=r.score).alignment = CENTER
        ws.cell(row=row, column=5, value=r.category)
        ws.merge_cells(start_row=row, start_column=5, end_row=row, end_column=7)


REGISTER_COLUMNS = [
    ("Risk ID", 9), ("Title", 32), ("Category", 22), ("Severity", 11), ("Score (L×I)", 10),
    ("Likelihood", 10), ("Impact", 9), ("Description", 50), ("Evidence (quote from document)", 50),
    ("Evidence verified", 11), ("Source document", 24), ("Source section", 24),
    ("Affected requirements", 16), ("Mitigation", 45), ("Open questions", 40), ("Knowledge-base reference", 28),
    ("Occurrences", 11), ("Owner", 14), ("Status", 11),
]


def _register(ws: Worksheet, result: AnalysisResult) -> None:
    ws.title = "Risk Register"
    _header(ws, 1, [c for c, _ in REGISTER_COLUMNS], [w for _, w in REGISTER_COLUMNS])
    ws.row_dimensions[1].height = 32
    for i, r in enumerate(result.risks, 2):
        values = [
            r.risk_id, r.title, r.category, r.severity, r.score, r.likelihood, r.impact,
            r.description, r.evidence, "Yes" if r.evidence_verified else "No – check",
            r.source_document, r.source_section, ", ".join(r.affected_requirements),
            r.mitigation, "\n".join(f"• {q}" for q in r.open_questions), r.reference, r.occurrences, "", "Open",
        ]
        for col, v in enumerate(values, 1):
            c = ws.cell(row=i, column=col, value=v)
            c.alignment, c.border = WRAP_TOP, BORDER
        _sev_cell(ws.cell(row=i, column=4), r.severity)
        for col in (1, 5, 6, 7, 10, 17, 19):
            ws.cell(row=i, column=col).alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
        if not r.evidence_verified:
            ws.cell(row=i, column=10).font = Font(color="C00000", bold=True)
    ws.freeze_panes = "C2"
    if result.risks:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(REGISTER_COLUMNS))}{len(result.risks) + 1}"


def _questions(ws: Worksheet, result: AnalysisResult) -> None:
    ws.title = "Open Questions"
    _header(ws, 1, ["#", "Risk ID", "Severity", "Question", "Answer", "Answered by"], [6, 10, 11, 70, 50, 16])
    row = 1
    for r in result.risks:
        for q in r.open_questions:
            row += 1
            vals = [row - 1, r.risk_id, r.severity, q, "", ""]
            for col, v in enumerate(vals, 1):
                c = ws.cell(row=row, column=col, value=v)
                c.alignment, c.border = WRAP_TOP, BORDER
            _sev_cell(ws.cell(row=row, column=3), r.severity)
    ws.freeze_panes = "A2"


def _run_info(ws: Worksheet, result: AnalysisResult) -> None:
    ws.title = "Run Info"
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 100
    rows = [
        ("LLM provider", result.provider),
        ("Model requested", result.model),
        ("Models used (calls)", ", ".join(f"{m} ({n})" for m, n in result.models_used.items())),
        ("Model switches", "\n".join(result.model_switches) or "none"),
        ("Context window (tokens)", result.context_tokens),
        ("Max output tokens / call", result.max_output_tokens),
        ("Chunk size (tokens)", result.chunk_tokens),
        ("LLM calls", result.llm_calls),
        ("Tokens (input / output)", f"{result.prompt_tokens:,} / {result.completion_tokens:,}"),
        ("Cost (USD, from OpenRouter)", result.cost_usd),
        ("RAG (retrieval)", result.rag_status),
        ("Embedding model", result.embedding_model or "–"),
        ("Pieces indexed / embedding calls", f"{result.rag_pieces_indexed} / {result.embedding_calls}"),
        ("Knowledge-base files", result.knowledge_base_files),
        ("Risks found / after de-duplication", f"{result.raw_risks} / {len(result.risks)}"),
        ("Started", result.started_at),
        ("Finished", result.finished_at),
        ("Chunks analysed", result.chunks_total),
        ("Chunks failed", len(result.chunks_failed)),
        ("Scoring", "Score = Likelihood × Impact (1–25). Critical ≥17, High 10–16, Medium 5–9, Low 1–4."),
        ("Evidence verified", "'No – check' means the quoted evidence could not be matched to the source text; review before relying on it."),
        ("Disclaimer", "AI-generated assessment. Review with the solution architect before use."),
    ]
    for i, (k, v) in enumerate(rows, 1):
        ws.cell(row=i, column=1, value=k).font = BOLD
        ws.cell(row=i, column=2, value=v).alignment = WRAP_TOP
    if result.chunks_failed:
        start = len(rows) + 2
        ws.cell(row=start, column=1, value="Failed chunks").font = BOLD
        for j, f in enumerate(result.chunks_failed, start):
            ws.cell(row=j, column=2, value=f).alignment = WRAP_TOP


def build_workbook(result: AnalysisResult) -> Workbook:
    wb = Workbook()
    _summary(wb.active, result)
    _register(wb.create_sheet(), result)
    _questions(wb.create_sheet(), result)
    _run_info(wb.create_sheet(), result)
    return wb


def to_excel_bytes(result: AnalysisResult) -> bytes:
    buf = io.BytesIO()
    build_workbook(result).save(buf)
    return buf.getvalue()
