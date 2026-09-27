from pathlib import Path

from openpyxl import load_workbook

from app.config import Settings
from app.excel_report import to_excel_bytes
from app.llm import MockLLM, ModelInfo, extract_json, resolve_sizing
from app.models import RiskItem
from app.parsing import chunk_document, load_document
from app.pipeline import consolidate, evidence_in_text, run_analysis
from app.models import Risk

SAMPLE = Path(__file__).resolve().parents[1] / "samples" / "sample_brd.md"


def test_extract_json_handles_fences_and_chatter():
    assert extract_json('```json\n{"risks": []}\n```') == {"risks": []}
    assert extract_json('Sure! Here it is: {"risks": [{"title": "x",}]} hope it helps') == {"risks": [{"title": "x"}]}
    assert extract_json('<think>hmm</think>[{"title": "y"}]') == {"risks": [{"title": "y"}]}


def test_risk_item_coercion():
    r = RiskItem(title="t", category="security", likelihood="High", impact="4/5", open_questions="a; b")
    assert (r.category, r.likelihood, r.impact, r.open_questions) == ("Security & Compliance", 4, 4, ["a", "b"])


def test_evidence_grounding():
    text = "NFR-02: Availability of 99.99% is required.\nOther text."
    assert evidence_in_text("Availability of 99.99% is required", text)
    assert evidence_in_text("availability of 99.99%  is required.", text)
    assert not evidence_in_text("The system must encrypt all data at rest", text)


def test_consolidate_merges_similar_titles():
    a = Risk(title="Third-party payment gateway dependency", category="Integration & Dependencies", likelihood=2, impact=4)
    b = Risk(title="Third party payment gateway dependency risk", category="Integration & Dependencies", likelihood=4, impact=3)
    c = Risk(title="PII handling", category="Security & Compliance")
    merged = consolidate([a, b, c])
    assert len(merged) == 2
    assert (merged[0].likelihood, merged[0].impact, merged[0].occurrences) == (4, 4, 2)


def test_resolve_sizing_auto_and_overrides():
    info = ModelInfo(id="m", context_length=1_000_000, max_completion_tokens=128_000)
    auto = resolve_sizing(Settings(llm_context_tokens=0, llm_max_output_tokens=0, chunk_tokens=0, max_chunk_tokens=30_000), info)
    assert (auto.context_tokens, auto.max_output_tokens, auto.chunk_tokens) == (1_000_000, 8_192, 30_000)

    small = ModelInfo(id="m", context_length=16_000, max_completion_tokens=4_096)
    sz = resolve_sizing(Settings(llm_context_tokens=0, llm_max_output_tokens=0, chunk_tokens=0), small)
    assert sz.max_output_tokens == 4_096 and sz.chunk_tokens + sz.max_output_tokens < 16_000

    # explicit chunk that doesn't fit is clamped
    sz = resolve_sizing(Settings(llm_context_tokens=0, llm_max_output_tokens=2_000, chunk_tokens=50_000), small)
    assert sz.chunk_tokens <= 16_000 - 2_000


def test_chunking_respects_size():
    doc = load_document("big.md", ("# Section\n" + "word " * 3000).encode(), "SRD")
    chunks = chunk_document(doc, 2000, 200)
    assert len(chunks) > 1 and all(len(c.text) <= 2000 for c in chunks)


def test_end_to_end_with_mock_llm():
    doc = load_document(SAMPLE.name, SAMPLE.read_bytes(), "BRD")
    result = run_analysis([doc], "Loyalty", cfg=Settings(llm_provider="mock"), llm=MockLLM())
    titles = {r.title for r in result.risks}
    assert "Data migration risk" in titles and "Sensitive data handling and compliance" in titles
    assert all(r.evidence_verified for r in result.risks)
    assert result.risks[0].risk_id == "TR-001"

    wb = load_workbook(__import__("io").BytesIO(to_excel_bytes(result)))
    assert wb.sheetnames == ["Summary", "Risk Register", "Open Questions", "Run Info"]
    assert wb["Risk Register"].max_row == len(result.risks) + 1


def test_consolidate_semantic_merges_across_categories_and_keeps_highest_score():
    a = Risk(title="Nightly batch conflicts with real-time visibility", category="Requirements Quality", likelihood=5, impact=4)
    b = Risk(title="Point accrual not real-time", category="Performance & Availability", likelihood=3, impact=3)
    c = Risk(title="PII without encryption", category="Security & Compliance", likelihood=3, impact=5)
    vectors = [[1.0, 0.1, 0.0], [0.9, 0.2, 0.0], [0.0, 0.2, 1.0]]
    merged = consolidate([b, c, a], vectors=[vectors[1], vectors[2], vectors[0]], semantic_threshold=0.65)
    assert len(merged) == 2
    top = merged[0]
    assert top.title == a.title and top.category == "Requirements Quality" and top.occurrences == 2
    # without vectors the differently-worded risks stay separate
    assert len(consolidate([a, b, c])) == 3
